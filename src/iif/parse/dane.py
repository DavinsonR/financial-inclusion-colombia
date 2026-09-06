"""Parsers de los anexos XLSX del DANE a tablas tidy (Parquet en data/interim/dane/).

Cada parser lee un libro tal como lo publica el DANE (encabezados en filas variables, notas al pie,
años con sufijos `p`/`pr`) y devuelve un DataFrame validado con pandera. Convenciones:
- códigos DIVIPOLA como texto con ceros a la izquierda (`dpto_ccdgo` 2 dígitos, `mpio_ccdgo` 5);
- `anio` entero y `estado_dato` ∈ {definitivo, provisional, preliminar} derivado del sufijo del año;
- valores monetarios en la unidad del anexo (miles de millones de pesos corrientes o constantes 2015);
- el agregado nacional se conserva con `dpto_ccdgo = '00'` y `ITAED` marca el residuo como `'RESTO'`.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd
import pandera.pandas as pa

from iif import config

INTERIM = config.DATA_INTERIM / "dane"
ESTADO = {"": "definitivo", "p": "provisional", "pr": "preliminar"}


# ----------------------------------------------------------------------------- utilidades
def _sheet_rows(path: Path, sheet: str) -> list[tuple]:
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl avisa por extensiones desconocidas del DANE
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
        finally:
            wb.close()


def _sheet_names(path: Path) -> list[str]:
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return wb.sheetnames
        finally:
            wb.close()


def _find_row(rows: list[tuple], startswith: str, col: int = 0, start: int = 0) -> int:
    for i in range(start, len(rows)):
        v = rows[i][col] if col < len(rows[i]) else None
        if isinstance(v, str) and v.strip().lower().startswith(startswith.lower()):
            return i
    raise ValueError(f"no se encontró una fila que empiece por {startswith!r} en la columna {col}")


def parse_year(v) -> tuple[int, str] | None:
    """'2024p' → (2024, 'provisional'); 2019 → (2019, 'definitivo'); otra cosa → None."""
    if v is None:
        return None
    m = re.fullmatch(r"\s*(\d{4})\s*(p|pr)?\s*", str(v))
    if not m:
        return None
    return int(m.group(1)), ESTADO[m.group(2) or ""]


def _num(v) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def normalize_name(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", " ", s.upper()).strip()


def _code(v, width: int) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.zfill(width) if s.isdigit() else None


def _year_columns(header: tuple) -> list[tuple[int, int, str]]:
    """Índices de columna con año en un encabezado; corta en la primera columna vacía tras el primer año."""
    out = []
    started = False
    for j, v in enumerate(header):
        y = parse_year(v)
        if y:
            out.append((j, *y))
            started = True
        elif started and v is None:
            break
    return out


# ----------------------------------------------------------------------------- PIB departamental total
def parse_pib_departamento(path: Path) -> pd.DataFrame:
    """`anex-PIBDep-TotalDep-*.xlsx`: Cuadro 1 corriente, Cuadro 2 constante 2015, Cuadro 3 per cápita."""
    measures = {
        "Cuadro 1": "pib_corriente_mm",
        "Cuadro 2": "pib_constante_2015_mm",
        "Cuadro 3": "pib_per_capita_corriente",
    }
    frames = []
    for sheet, medida in measures.items():
        rows = _sheet_rows(path, sheet)
        h = _find_row(rows, "Código Departamento")
        years = _year_columns(rows[h])
        for r in rows[h + 1 :]:
            name = r[1]
            if not isinstance(name, str) or not name.strip():
                if r[0] is not None and isinstance(r[0], str) and r[0].lower().startswith("fuente"):
                    break
                continue
            code = _code(r[0], 2) or ("00" if normalize_name(name) == "COLOMBIA" else None)
            if code is None:
                continue
            for j, anio, estado in years:
                frames.append((code, name.strip(), anio, estado, medida, _num(r[j])))
    long = pd.DataFrame(
        frames, columns=["dpto_ccdgo", "departamento", "anio", "estado_dato", "medida", "valor"]
    )
    wide = _widen(long, ["dpto_ccdgo", "departamento", "anio", "estado_dato"])
    return PIB_DEPARTAMENTO.validate(wide.sort_values(["dpto_ccdgo", "anio"]).reset_index(drop=True))


PIB_DEPARTAMENTO = pa.DataFrameSchema(
    {
        "dpto_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{2}$")),
        "departamento": pa.Column(str),
        "anio": pa.Column(int, pa.Check.in_range(2000, 2035)),
        "estado_dato": pa.Column(str, pa.Check.isin(list(ESTADO.values()))),
        "pib_corriente_mm": pa.Column(float, pa.Check.gt(0), nullable=True),
        "pib_constante_2015_mm": pa.Column(float, pa.Check.gt(0), nullable=True),
        "pib_per_capita_corriente": pa.Column(float, pa.Check.gt(0), nullable=True),
    },
    unique=["dpto_ccdgo", "anio"],
    coerce=True,
)


def _widen(long: pd.DataFrame, index: list[str]) -> pd.DataFrame:
    """Una columna por `medida` sobre las combinaciones observadas. Nunca un pivote con `dropna=False`:
    construye el producto cartesiano de todos los niveles del índice y agota la memoria (B-026)."""
    wide = long.set_index([*index, "medida"])["valor"].unstack("medida").reset_index()
    return wide.rename_axis(None, axis=1)


# ----------------------------------------------------------------------------- VA departamental por actividad
def parse_va_departamento_actividad(path: Path, dept_codes: dict[str, str]) -> pd.DataFrame:
    """`anex-PIBDep-departamento-*.xlsx`: un bloque por departamento y actividad (Cuadro 1 corriente, Cuadro 2 constante)."""
    measures = {"Cuadro 1": "va_corriente_mm", "Cuadro 2": "va_constante_2015_mm"}
    recs = []
    for sheet, medida in measures.items():
        rows = _sheet_rows(path, sheet)
        i = 0
        while i < len(rows):
            v = rows[i][0]
            # Título de bloque: "<Departamento>: valor agregado por actividad económica ..."
            if isinstance(v, str) and ":" in v and normalize_name(v.split(":", 1)[0]) in dept_codes:
                dept = v.split(":", 1)[0].strip()
                h = _find_row(rows, "Clasificación", start=i)
                years = _year_columns(rows[h])
                k = h + 1
                while k < len(rows) and rows[k][2] is not None:
                    r = rows[k]
                    cuenta = str(r[0]).strip() if r[0] is not None else None
                    seccion = str(r[1]).strip() if r[1] is not None else None
                    actividad = str(r[2]).strip()
                    for j, anio, estado in years:
                        recs.append((dept, cuenta, seccion, actividad, anio, estado, medida, _num(r[j])))
                    k += 1
                i = k
            i += 1
    long = pd.DataFrame(
        recs,
        columns=[
            "departamento",
            "cuenta",
            "seccion_ciiu",
            "actividad",
            "anio",
            "estado_dato",
            "medida",
            "valor",
        ],
    )
    long["dpto_ccdgo"] = long["departamento"].map(lambda s: dept_codes[normalize_name(s)])
    n_dep = long["dpto_ccdgo"].nunique()
    if n_dep != 33:
        raise ValueError(f"se esperaban 33 bloques departamentales y se leyeron {n_dep}")
    for c in ("cuenta", "seccion_ciiu"):
        long[c] = long[c].fillna("")  # el índice del unstack no admite nulos
    index = ["dpto_ccdgo", "departamento", "cuenta", "seccion_ciiu", "actividad", "anio", "estado_dato"]
    wide = _widen(long, index)
    return VA_DEPARTAMENTO_ACTIVIDAD.validate(
        wide.sort_values(["dpto_ccdgo", "anio", "cuenta", "seccion_ciiu"]).reset_index(drop=True)
    )


VA_DEPARTAMENTO_ACTIVIDAD = pa.DataFrameSchema(
    {
        "dpto_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{2}$")),
        "departamento": pa.Column(str),
        "cuenta": pa.Column(str),
        "seccion_ciiu": pa.Column(str),
        "actividad": pa.Column(str),
        "anio": pa.Column(int, pa.Check.in_range(2000, 2035)),
        "estado_dato": pa.Column(str),
        "va_corriente_mm": pa.Column(float, nullable=True),
        "va_constante_2015_mm": pa.Column(float, nullable=True),
    },
    unique=["dpto_ccdgo", "actividad", "anio"],
    coerce=True,
)


# ----------------------------------------------------------------------------- VA municipal
def parse_va_municipio(path: Path) -> pd.DataFrame:
    """`anex-PIBDep-ValorAgreMuni-*.xlsx`: Cuadro 1 serie total; Cuadros 2+ un año cada uno con actividades y peso."""
    rows = _sheet_rows(path, "Cuadro 1")
    h = _find_row(rows, "Código Municipio")
    years = _year_columns(rows[h])
    recs = []
    for r in rows[h + 1 :]:
        code = _code(r[0], 5)
        if code is None:
            if isinstance(r[0], str) and r[0].lower().startswith("fuente"):
                break
            continue
        for j, anio, estado in years:
            recs.append(
                (code, str(r[1]).strip(), _code(r[2], 2), str(r[3]).strip(), anio, estado, _num(r[j]))
            )
    serie = pd.DataFrame(
        recs,
        columns=[
            "mpio_ccdgo",
            "municipio",
            "dpto_ccdgo",
            "departamento",
            "anio",
            "estado_dato",
            "va_corriente_mm",
        ],
    )
    # Cuadros por año: el índice del libro dice qué año es cada cuadro.
    idx = _sheet_rows(path, "Índice")
    year_of_sheet = {}
    for r in idx:
        if (
            isinstance(r[0], str)
            and r[0].startswith("Cuadro")
            and isinstance(r[1], str)
            and r[1].startswith("Año")
        ):
            y = parse_year(r[1].replace("Año", ""))
            if y:
                year_of_sheet[r[0].strip()] = y[0]
    parts = []
    for sheet, anio in year_of_sheet.items():
        rs = _sheet_rows(path, sheet)
        hh = _find_row(rs, "Código Municipio")
        for r in rs[hh + 1 :]:
            code = _code(r[0], 5)
            if code is None:
                continue
            parts.append((code, anio, _num(r[4]), _num(r[5]), _num(r[6]), _num(r[7]), _num(r[8])))
    act = pd.DataFrame(
        parts,
        columns=[
            "mpio_ccdgo",
            "anio",
            "va_primarias_mm",
            "va_secundarias_mm",
            "va_terciarias_mm",
            "va_total_cuadro_mm",
            "peso_relativo_pct",
        ],
    )
    out = serie.merge(act, on=["mpio_ccdgo", "anio"], how="left", validate="one_to_one")
    return VA_MUNICIPIO.validate(out.sort_values(["mpio_ccdgo", "anio"]).reset_index(drop=True))


VA_MUNICIPIO = pa.DataFrameSchema(
    {
        "mpio_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{5}$")),
        "municipio": pa.Column(str),
        "dpto_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{2}$")),
        "departamento": pa.Column(str),
        "anio": pa.Column(int, pa.Check.in_range(2005, 2035)),
        "estado_dato": pa.Column(str),
        "va_corriente_mm": pa.Column(float, pa.Check.ge(0), nullable=True),
        "va_primarias_mm": pa.Column(float, nullable=True),
        "va_secundarias_mm": pa.Column(float, nullable=True),
        "va_terciarias_mm": pa.Column(float, nullable=True),
        "va_total_cuadro_mm": pa.Column(float, nullable=True),
        "peso_relativo_pct": pa.Column(float, pa.Check.in_range(0, 100), nullable=True),
    },
    checks=[
        pa.Check(
            lambda df: (df["mpio_ccdgo"].str[:2] == df["dpto_ccdgo"]).all(),
            error="los dos primeros dígitos del municipio deben ser el departamento",
        )
    ],
    unique=["mpio_ccdgo", "anio"],
    coerce=True,
)


# ----------------------------------------------------------------------------- población
AREA = {"cabecera municipal": "cabecera", "centros poblados y rural disperso": "resto", "total": "total"}


def _parse_poblacion(path: Path, sheet: str, fuente: str) -> pd.DataFrame:
    rows = _sheet_rows(path, sheet)
    h = _find_row(rows, "DP")
    header = [str(v).strip().upper() if v is not None else "" for v in rows[h]]
    col = {name: j for j, name in enumerate(header)}
    recs = []
    for r in rows[h + 1 :]:
        dp = _code(r[col["DP"]], 2)
        if dp is None:
            continue
        # En los libros retro el orden DPMP/MPIO cambia: el código es la celda de 5 dígitos.
        mpio = mpio_name = None
        if "MPIO" in col:
            a, b = r[col["MPIO"]], r[col.get("DPMP", col["MPIO"])]
            mpio = _code(a, 5) or _code(b, 5)
            mpio_name = str(b if _code(a, 5) else a).strip()
        anio = parse_year(r[col["AÑO"]])
        area_raw = str(r[col["ÁREA GEOGRÁFICA"]]).strip().lower()
        total_col = col.get("TOTAL", col.get("POBLACIÓN"))
        recs.append(
            (
                dp,
                str(r[col["DPNOM"]]).strip(),
                mpio,
                mpio_name,
                anio[0] if anio else None,
                AREA.get(area_raw, area_raw),
                _num(r[total_col]),
                fuente,
            )
        )
    df = pd.DataFrame(
        recs,
        columns=[
            "dpto_ccdgo",
            "departamento",
            "mpio_ccdgo",
            "municipio",
            "anio",
            "area",
            "poblacion",
            "fuente",
        ],
    )
    return df


def parse_poblacion_departamento(path: Path) -> pd.DataFrame:
    df = _parse_poblacion(path, "PobDepartamentalxÁrea", "proyeccion_2018_2050").drop(
        columns=["mpio_ccdgo", "municipio"]
    )
    return POBLACION_DEPARTAMENTO.validate(df)


def parse_poblacion_municipio(path_2018: Path, path_retro: Path | None = None) -> pd.DataFrame:
    parts = [_parse_poblacion(path_2018, "PobMunicipalxÁrea", "proyeccion_2018_2042")]
    if path_retro is not None:
        parts.append(_parse_poblacion(path_retro, _sheet_names(path_retro)[0], "retroproyeccion_2005_2017"))
    df = pd.concat(parts, ignore_index=True)
    return POBLACION_MUNICIPIO.validate(df.sort_values(["mpio_ccdgo", "anio", "area"]).reset_index(drop=True))


POBLACION_DEPARTAMENTO = pa.DataFrameSchema(
    {
        "dpto_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{2}$")),
        "departamento": pa.Column(str),
        "anio": pa.Column(int, pa.Check.in_range(2000, 2060)),
        "area": pa.Column(str, pa.Check.isin(["cabecera", "resto", "total"])),
        "poblacion": pa.Column(float, pa.Check.ge(0)),
        "fuente": pa.Column(str),
    },
    unique=["dpto_ccdgo", "anio", "area"],
    coerce=True,
)

POBLACION_MUNICIPIO = pa.DataFrameSchema(
    {
        "dpto_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{2}$")),
        "departamento": pa.Column(str),
        "mpio_ccdgo": pa.Column(str, pa.Check.str_matches(r"^\d{5}$")),
        "municipio": pa.Column(str),
        "anio": pa.Column(int, pa.Check.in_range(2000, 2060)),
        "area": pa.Column(str, pa.Check.isin(["cabecera", "resto", "total"])),
        "poblacion": pa.Column(float, pa.Check.ge(0)),
        "fuente": pa.Column(str),
    },
    unique=["mpio_ccdgo", "anio", "area"],
    coerce=True,
)


# ----------------------------------------------------------------------------- ITAED
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4}


def parse_itaed(path: Path) -> pd.DataFrame:
    """`anex-ITAED-*.xlsx`, Cuadro 1, primer bloque: índice trimestral original (base 2015) por departamento."""
    rows = _sheet_rows(path, "Cuadro 1")
    h = _find_row(rows, "Código Departamento")
    year_row, q_row = rows[h], rows[h + 1]
    cols = []
    current = None
    for j in range(2, len(year_row)):
        y = parse_year(year_row[j])
        if y:
            current = y
        q = str(q_row[j]).strip() if q_row[j] is not None else None
        if current and q in ROMAN:
            cols.append((j, current[0], current[1], ROMAN[q]))
    recs = []
    for r in rows[h + 2 :]:
        name = r[1]
        if not isinstance(name, str) or not name.strip():
            if isinstance(r[0], str) and r[0].lower().startswith("fuente"):
                break
            continue
        code = _code(r[0], 2)
        if code is None:
            code = (
                "00"
                if normalize_name(name) == "COLOMBIA"
                else ("RESTO" if str(r[0]).strip() == "*" else None)
            )
        if code is None:
            continue
        for j, anio, estado, tri in cols:
            v = _num(r[j])
            if v is not None:
                recs.append((code, name.strip(), anio, tri, estado, v))
    df = pd.DataFrame(
        recs, columns=["dpto_ccdgo", "departamento", "anio", "trimestre", "estado_dato", "indice_original"]
    )
    return ITAED.validate(df)


ITAED = pa.DataFrameSchema(
    {
        "dpto_ccdgo": pa.Column(str, pa.Check.str_matches(r"^(\d{2}|RESTO)$")),
        "departamento": pa.Column(str),
        "anio": pa.Column(int, pa.Check.in_range(2010, 2035)),
        "trimestre": pa.Column(int, pa.Check.in_range(1, 4)),
        "estado_dato": pa.Column(str),
        "indice_original": pa.Column(float, pa.Check.gt(0)),
    },
    unique=["dpto_ccdgo", "anio", "trimestre"],
    coerce=True,
)


# ----------------------------------------------------------------------------- orquestación
def parse_all(manifest: Path | None = None, out_dir: Path | None = None) -> dict[str, Path]:
    """Lee la última descarga de cada fuente DANE según el manifiesto y escribe los Parquet tidy."""
    from iif.acquire.dane import latest_file

    out_dir = out_dir or INTERIM
    out_dir.mkdir(parents=True, exist_ok=True)
    written: dict[str, Path] = {}

    def need(sid: str) -> Path:
        p = latest_file(sid, manifest)
        if p is None or not p.exists():
            raise FileNotFoundError(f"{sid}: no hay descarga en el manifiesto; corre `iif acquire {sid}`")
        return p

    pib = parse_pib_departamento(need("dane-pib-total"))
    written["pib_departamento_anual"] = out_dir / "pib_departamento_anual.parquet"
    pib.to_parquet(written["pib_departamento_anual"], index=False)

    codes = {
        normalize_name(n): c
        for c, n in pib[["dpto_ccdgo", "departamento"]].drop_duplicates().itertuples(index=False)
    }
    va_dep = parse_va_departamento_actividad(need("dane-pib-departamento"), codes)
    written["va_departamento_actividad_anual"] = out_dir / "va_departamento_actividad_anual.parquet"
    va_dep.to_parquet(written["va_departamento_actividad_anual"], index=False)

    va_mun = parse_va_municipio(need("dane-va-municipio"))
    written["va_municipio_anual"] = out_dir / "va_municipio_anual.parquet"
    va_mun.to_parquet(written["va_municipio_anual"], index=False)

    pob_dep = parse_poblacion_departamento(need("dane-poblacion-departamental"))
    written["poblacion_departamento_anual"] = out_dir / "poblacion_departamento_anual.parquet"
    pob_dep.to_parquet(written["poblacion_departamento_anual"], index=False)

    retro = latest_file("dane-poblacion-municipal-retro", manifest)
    pob_mun = parse_poblacion_municipio(
        need("dane-poblacion-municipal"), retro if retro and retro.exists() else None
    )
    written["poblacion_municipio_anual"] = out_dir / "poblacion_municipio_anual.parquet"
    pob_mun.to_parquet(written["poblacion_municipio_anual"], index=False)

    itaed = parse_itaed(need("dane-itaed"))
    written["itaed_departamento_trimestre"] = out_dir / "itaed_departamento_trimestre.parquet"
    itaed.to_parquet(written["itaed_departamento_trimestre"], index=False)
    return written
